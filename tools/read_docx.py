#!/usr/bin/env python
# -*- coding: utf-8 -*-
r"""read_docx.py —— 取 .docx 全文（段落 + 表格），只读、不改。 v1.0.0
用法: read_docx.py <文件> [--head N] [--grep 关键词]
"""
from __future__ import annotations
import argparse
import pathlib
import sys

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8")
    except Exception:
        pass


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("path")
    ap.add_argument("--head", type=int, default=0)
    ap.add_argument("--grep", default="")
    a = ap.parse_args()
    p = pathlib.Path(a.path)
    if not p.is_file():
        print("[ERR] 不存在：%s" % p); return 2
    try:
        import docx  # python-docx
    except Exception as e:
        print("[ERR] 缺 python-docx：%s" % e); return 3
    d = docx.Document(str(p))
    print("═══ %s（%d B）═══" % (p.name, p.stat().st_size))
    print("  段落 %d ｜ 表格 %d" % (len(d.paragraphs), len(d.tables)))
    print()
    n = 0
    for para in d.paragraphs:
        t = para.text.strip()
        if not t:
            continue
        if a.grep and a.grep not in t:
            continue
        n += 1
        if a.head and n > a.head:
            print("  …（其余略，共 %d 段）" % len([x for x in d.paragraphs if x.text.strip()]))
            break
        print("  %3d| %s" % (n, t[:150]))
    for ti, tb in enumerate(d.tables, 1):
        print()
        print("  ── 表 %d（%d 行 × %d 列）──" % (ti, len(tb.rows), len(tb.columns)))
        for ri, row in enumerate(tb.rows):
            if a.head and ri >= a.head:
                print("     …（其余略）"); break
            cells = [c.text.strip().replace("\n", " ")[:44] for c in row.cells]
            print("     | " + " | ".join(cells) + " |")
    return 0


if __name__ == "__main__":
    sys.exit(main())
